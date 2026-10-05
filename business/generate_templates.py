"""Generate blank business document templates; never put customer data in this file."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parent
INK, GREEN, PAPER, RULE = '111111', '526657', 'F4F4F0', 'C7C7BE'
VERSION = '2026-10-05'

CONTRACT = [
 ('第1条　目的・契約書類', [
  '発注者（以下「甲」）と、個人事業として専門サイト工房を運営する受注者（以下「乙」）は、甲の事業用Webサイトの制作について本契約を締結します。契約の対象は、別紙1で特定した静的Webサイトです。',
  '本契約書、別紙1「案件仕様・合意事項」、別紙2「見積書」が契約内容を構成します。見積書の番号・版・発行日と、添付した実際のファイルを一致させます。記載が食い違う場合は制作前に確認し、双方が文書で合意した個別条件を優先します。未記入の事項は双方で確定してから契約します。無料相談・制作体験の操作や問い合わせ送信だけで有料契約は成立しません。',
 ]),
 ('第2条　制作範囲・修正', [
  '制作ページ・セクション・機能・使用サービス・対応環境・納品方法は別紙1に限定して定めます。標準範囲は、構成の整理、デザイン、スマホ対応、公開設定、検索向けの基本設定、お問い合わせフォーム1つ、地図、外部予約サービス・SNSへのリンクです。原稿を甲が支給するプライバシーポリシー1ページは、プランの制作ページ数とは別に料金内で掲載します。その他の法務ページは別途合意します。',
  '確認案への軽微な修正は2回まで制作費に含みます。甲がまとめた要望への対応を1回と数え、合意した構成・機能の範囲で文章・写真・配色等を調整します。乙の制作ミスや契約内容との不一致を直す作業は、この回数に含めません。',
  'CMS、決済・会員・検索システム、独自の予約機能、原稿の新規作成、撮影、追加の法務ページ、大幅な構成・デザイン変更等は、別紙に明記した場合を除き含みません。制作体験の生成結果は相談のたたき台であり、提案された機能すべてを基本料金で実装する約束ではありません。',
 ]),
 ('第3条　素材・アカウント・外部サービス', [
  '甲は合意した期限までに、原稿・画像・事業情報・利用可能な素材を支給し、その内容の正確性と必要な利用許諾を確認します。乙も自ら調達する素材の利用条件を確認し、双方は権利上の問題を知ったときは速やかに相手へ連絡します。',
  'ドメイン・公開環境・フォーム受信先等の管理主体と必要な操作は別紙1で定めます。原則として甲が所有・管理するアカウントを利用し、乙への作業権限は必要な期間・範囲に限定します。フォーム受信承認・ドメイン取得等、甲が行う操作は甲が期限までに実施します。パスワード等の共有方法は別途安全な方法を合意します。',
  '第三者のサービス・素材・フォント等には各提供者の利用条件が適用されます。費用・個人情報の送信先・甲に必要な操作を公開前に確認します。外部サービスの障害・仕様変更・費用変更等が分かった場合は、双方で対応範囲・費用・納期を確認します。乙の設定ミス等まで外部サービスを理由に免責するものではありません。',
 ]),
 ('第4条　制作開始・納期・変更', [
  '制作範囲の確定、必要素材の受領、着手金の入金確認が揃った日を制作開始日とします。制作期間は開始から1週間（7日間）を目安とし、実際の最終確認案提出日・公開予定日を別紙1で合意します。目安だけを確定納期として扱いません。',
  '確認回答・素材・アカウント操作の遅れや追加変更等がある場合、乙は影響を速やかに説明し、費用・納期の変更を双方がメール等の文書で合意してから進めます。合意のない追加料金を請求しません。災害等の不可抗力についても速やかに連絡し、日程・対応を協議します。',
 ]),
 ('第5条　制作費・支払い・外部費用', [
  '制作費の税込総額、消費税、工程別内訳、追加費用、支払期限、振込先は別紙2で定めます。甲は制作着手前に制作費の50％を、最終確認後・公開前に残り50％を、乙が示す口座へ銀行振込で支払います。甲が最終内容を確認できる状態で提示してから残金を請求します。',
  'ドメイン・サーバー・有料素材・有料サービス等の外部実費は制作費と分け、項目・費用・負担者を事前に文書で合意します。銀行振込手数料は甲の負担とします。追加作業・追加修正は実施前に見積り、双方の承諾を得ます。',
  '任意の月額更新サポートは本契約に自動的に含まれず、希望する場合に別の申込書で開始日・料金・条件を合意します。',
 ]),
 ('第6条　確認・納品・公開', [
  '乙は別紙1の確認環境で完成内容を提示し、甲は原則として提示後5営業日以内に、承認または具体的な修正事項をメール等で回答します。回答がないことだけで合格・承認とは扱わず、乙は連絡して確認日程を調整します。',
  '契約内容との不一致は乙が修正して再確認を依頼します。甲の最終承認と残金入金の確認後、乙は合意した公開先への反映・納品を行います。納品物は制作したHTML・CSS・JavaScript・使用可能な画像等と更新案内で、具体的なファイル形式・アカウントの引渡し方法は別紙1に定めます。',
  '公開時にフォーム送信・リンク等の動作を双方で確認します。検索向けの基本設定は行いますが、検索順位・売上・問い合わせ件数を保証するものではありません。',
 ]),
 ('第7条　権利・実績掲載', [
  '乙が本案件のために新たに作成し、著作権が発生して乙に帰属する成果物の著作権は、制作費全額の支払いと納品が完了したときに甲へ移転します。この移転には著作権法第27条・第28条の権利を含みます。乙は甲及び甲が認めた利用者の契約目的に沿う利用・改変について著作者人格権を行使しません。',
  '甲の支給物、第三者の素材・フォント・ライブラリ等、乙が以前から保有する汎用コード・テンプレート・制作手法は移転対象から除きます。乙の保有物については、甲が納品サイトを運営・改修するための非独占的な利用を認めます。第三者素材はそのライセンス条件に従い、移転対象外の具体的なものを納品時に案内します。',
  'サイト名・画面・制作過程等を乙の実績として公開する場合は、公開内容・時期を示して甲の事前承諾を得ます。別紙の実績掲載欄は「許可しない」が初期状態です。',
 ]),
 ('第8条　不具合・契約内容との不一致', [
  '乙の制作ミスや合意した仕様との不一致は、お客様都合の修正やキャンセルと区別します。甲は内容を乙へ連絡し、乙は原因を確認して修正等の対応を行います。公開後30日以内の連絡を初期確認期間としますが、この期間を過ぎたことだけを理由に法令上の権利を失わせるものではありません。',
  '甲・第三者による改変や、公開後の外部サービス変更等により新たに生じた問題は、原因と対応範囲を確認して別途協議します。契約不履行等の修正・解除・損害賠償は、契約内容と適用される法令に従って扱います。',
 ]),
 ('第9条　キャンセル・精算・解除', [
  '契約前の相談・見積りは無料です。甲が契約後、制作着手前にキャンセルする場合、乙は受領金を原則返金し、甲と事前に合意して発生した外部実費があるときだけ精算します。着手金を一律に返金不可とはしません。',
  '甲の都合で制作着手後にキャンセルする場合、別紙2の工程別金額と実施済み作業の記録に基づき、実施分と事前合意済みの外部実費を精算します。未実施の工程全額や未承諾の追加費用を請求しません。乙は計算根拠を甲へ提示し、受領金に余剰があれば精算額合意後10営業日以内に返金します。不足がある場合は差額と支払期限を文書で確認します。',
  '一方に契約違反がある場合、相手方は内容を示し、原則として合理的な期間を定めて是正を求めます。改善されない場合等は、契約内容と法令に従って解除・精算等を行います。乙の不履行による解除を甲都合のキャンセルと同じ条件で扱わず、甲の法令上の請求権を制限しません。',
  'キャンセル時の成果物の引渡し・利用範囲は、支払対象となった作業と第三者の権利を踏まえて文書で合意します。',
 ]),
 ('第10条　秘密情報・個人情報', [
  '双方は、本件に関連して知った未公開の事業情報、連絡先、認証情報等を、履行に必要な範囲を超えて利用・開示しません。公知情報、正当に取得した情報、法令に基づく開示は除きます。この義務は契約終了後も存続します。',
  '乙が扱う個人情報は本件の連絡・制作・請求・サポート等に必要な範囲に限り、適切に管理します。甲のサイトから利用者情報を送信するサービス・送信先・運用担当者は別紙1で確認し、甲は事業内容に合うプライバシー原稿・運用方針を用意します。第三者への新たな送信や処理が必要な場合、乙は甲へ説明して合意を得ます。',
  '契約終了後、不要となったアクセス権・作業用データは整理・返却・削除します。法令上の保存義務や取引の記録として必要な書類は、利用目的と必要性に応じて保管します。甲のフォーム受信メールの継続管理は、別途委託した場合を除き甲が行います。',
 ]),
 ('第11条　責任・準拠法・協議', [
  '双方は、自己の責めに帰すべき事由で相手に損害を与えた場合、適用される法令に従って責任を負います。原因・影響・対応方法を速やかに共有し、被害の拡大を防ぐために協力します。',
  '本契約には日本法を適用します。定めのない事項や解釈に疑義がある事項は、双方が誠実に協議します。協議で解決しない場合の裁判所は法令上の管轄に従います。',
 ]),
 ('第12条　契約の成立・保管', [
  '本契約と別紙の確定版について、双方が署名するか、電子署名等により承諾したときに契約が成立します。メールで締結する場合は、確定した同一のPDF一式・版番号を特定し、双方が本契約及び別紙すべてに承諾する旨を明示して、そのメールと添付書類を保管します。',
  '正式な氏名・住所・連絡可能な電話番号は、契約当事者欄に記入して相手方へ提供します。インターネットへ一般公開しない運用と、契約相手への提供は別です。',
 ]),
]

SUPPORT = [
 ('1　対象・対応範囲', [
  '制作契約とは別に、指定サイトの既存ページの文章・写真の差し替え等、軽微な更新を行う任意のサービスです。月額5,500円（税込）、月2回・合計60分まで。まとめて受けた依頼への対応を1回とし、実作業の合計時間を上限とします。回数・時間の翌月への繰越はありません。',
  '新規ページ・追加機能・大幅なデザイン変更・上限を超える作業は含みません。受注者は上限を超えると分かった時点で作業を止め、別途見積り・承諾を得てから追加対応します。制作契約の不一致や制作ミスを直す対応は、この更新枠を消費しません。',
 ]),
 ('2　期間・支払い・解約', [
  '契約は暦月単位（毎月1日〜末日）とし、記載した開始月から適用します。当月分は月初に銀行振込で支払い、具体的な期限は請求書に記載します。開始月の支払期限は下欄にも記入します。振込手数料は発注者の負担です。',
  '当月末までに受注者の窓口へメールで解約を申し出ると翌月から終了します。受注者は受付をメールで返信します。申出後に返信が確認できない場合は再連絡してください。受注者から終了する場合は原則として前月末までに連絡し、受領済みで対応できない期間・作業分がある場合は協議して精算します。',
  '日割り・途中月開始・特別な更新期間を希望する場合は、料金・回数・支払期限・解約期限を事前に個別合意欄で定めます。未利用の更新枠の繰越・返金はありません。ただし受注者の不履行等や法令上の請求は別に扱います。',
 ]),
 ('3　依頼・確認・管理', [
  '依頼は合意したメール窓口で受け付けます。支給素材と具体的な変更箇所を確認し、対応予定日を案内します。即日・24時間対応を約束するサービスではありません。依頼内容・作業時間・残り回数を双方が確認できるよう記録します。',
  'ドメイン・サーバー・有料サービス等の実費、第三者のサービス障害、監視・緊急復旧、データの常時バックアップ、集客・検索順位の保証は本更新サービスに含みません。必要な内容は別途合意します。',
  'アクセス権限は必要な範囲に限定し、双方は秘密情報・個人情報を適切に扱います。契約終了時の権限解除・作業用データの整理を行います。責任・準拠法・紛争の扱いは適用法令に従い、問題があれば双方で協議します。',
 ]),
 ('4　申込み・確定版の保管', [
  '本申込書の未記入欄と個別条件を確定し、双方が署名・電子署名するか、同じ確定版PDFを特定してメールで承諾したときに成立します。本サポートに申し込まなくても納品されたサイトは利用できます。',
 ]),
]

def field(document, label, value='［記入］'):
    table = document.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(4), Cm(12.2)
    table.cell(0,0).width, table.cell(0,1).width = Cm(4), Cm(12.2)
    trpr = table.rows[0]._tr.get_or_add_trPr()
    trpr.append(OxmlElement('w:cantSplit'))
    table.cell(0, 0).text, table.cell(0, 1).text = label, value
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            for run in p.runs:run.font.size = Pt(9)
    for cell in table.rows[0].cells[:1]:
        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), PAPER); cell._tc.get_or_add_tcPr().append(shading)

def document_base(title, caption):
    d = Document(); sec=d.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.top_margin,sec.bottom_margin,sec.left_margin,sec.right_margin=Cm(1.8),Cm(1.7),Cm(2),Cm(2)
    normal=d.styles['Normal']; normal.font.name='Noto Sans CJK JP'; normal.font.size=Pt(9.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK JP')
    normal.paragraph_format.line_spacing=1.25; normal.paragraph_format.space_after=Pt(5)
    for s,size in [('Title',24),('Heading 1',14),('Heading 2',11)]:
        style=d.styles[s]; style.font.name='Noto Sans CJK JP'; style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string(INK)
        style.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Sans CJK JP')
        style.paragraph_format.space_before=Pt(12);style.paragraph_format.space_after=Pt(6)
    title_border=d.styles['Title'].element.find(qn('w:pPr')).find(qn('w:pBdr'))
    for edge in title_border:
        edge.set(qn('w:color'),RULE)
        edge.attrib.pop(qn('w:themeColor'),None)
    d.styles['Caption'].font.color.rgb=RGBColor.from_string(GREEN)
    d.styles['Caption'].font.bold=False
    h=sec.header.paragraphs[0];h.text='専門サイト工房  /  pageatelier.jp';h.style=d.styles['Caption']
    f=sec.footer.paragraphs[0];f.alignment=2
    f.add_run('ひな形 ' + VERSION + '  |  ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
    d.add_paragraph(caption)
    d.add_heading(title,0)
    p=d.add_paragraph('記入用ひな形｜未記入のまま契約・送付に使用しないでください。')
    p.runs[0].font.color.rgb=RGBColor.from_string(GREEN)
    d.core_properties.author='専門サイト工房'
    d.core_properties.title=title
    return d

def write_contract():
    d=document_base('Webサイト制作契約書','01 / WEBSITE PRODUCTION')
    for label,value in [
      ('契約番号・版','［例：PA-2026-001 / 第1版］'),('契約日','［年・月・日］'),
      ('発注者（甲）','［法人の正式名称・代表者／個人事業主の正式氏名］'),
      ('受注者（乙）','［運営者の正式氏名］（個人事業・サービス名：専門サイト工房）'),
      ('案件名・対象サイト','［案件名・URLまたは公開予定ドメイン］'),
      ('制作費（税込）','［見積書の確定総額］円'),('添付見積書','［番号・版・発行日・確定PDFのファイル名］')]: field(d,label,value)
    for title,paragraphs in CONTRACT:
        d.add_heading(title,1)
        for paragraph in paragraphs:d.add_paragraph(paragraph)
    d.add_heading('契約当事者・署名欄',1)
    for party in ['甲（発注者）','乙（受注者）']:
        d.add_heading(party,2)
        for label,value in [('正式名称・氏名','［個人事業は正式氏名。法人は正式法人名・代表者名］'),('住所','［郵便番号・正確な住所］'),('電話番号','［連絡可能な電話番号］'),('メール','eyu.eyu0806@gmail.com' if party.startswith('乙') else '［連絡先メール］'),('署名／電子署名・日付','［署名等・年・月・日］')]:field(d,label,value)
    d.add_page_break();d.add_heading('別紙1　案件仕様・合意事項',0)
    d.add_paragraph('本契約と一緒に確定・保存します。該当しない欄は「なし」、未決事項は契約前に確定してください。')
    for label,value in [
      ('契約番号・版／見積番号','［本契約・見積書と同じ番号・版］'),
      ('事業内容・サイトの目的','［事業用であること、想定する訪問者・目的］'),
      ('制作プラン','［1ページ55,000円／5ページまで110,000円、いずれも税込］'),
      ('ページ・セクション','［全ページの名称・URL・主な内容／1ページの場合は8セクション以内の一覧］'),
      ('プライバシーページ','［原稿支給1ページを基本料金内・ページ数外で掲載。URL・原稿支給日を記入］'),
      ('その他の法務ページ','［なし／ページ名・支給原稿・別途費用］'),
      ('掲載量・支給素材','［原稿・画像点数等の目安、甲から支給する内容］'),
      ('フォーム','［1つ。項目・送信サービス・受信メール・テスト方法・受信承認担当］'),
      ('地図・外部リンク','［住所、掲載方法、予約URL・SNSのURL／なし］'),
      ('追加機能・対象外','［追加機能と見積項目／なし。CMS・決済・会員・検索・独自予約等の扱いを確認］'),
      ('修正・確認担当','軽微な修正2回。［甲の回答担当・要望の取りまとめ方法］'),
      ('対応環境','［例：現行のChrome・Safari・Edge、幅1440／1200／390px。案件に合わせ確定］'),
      ('素材・権限の準備期限','［年・月・日／甲に必要な操作］'),
      ('制作開始日','［仕様・必要素材・着手金が揃う予定日。実際の開始をメールで確認］'),
      ('最終確認案・公開予定日','［年・月・日］／［年・月・日］。開始から7日間目安、個別の確定日を記入。'),
      ('公開・管理するアカウント','［ドメイン・サーバー等、甲または乙の管理主体、作業権限、外部費用］'),
      ('納品・引渡し方法','［ZIP／甲のリポジトリ等。ソース・使用可能な画像・更新案内・ライセンス一覧］'),
      ('実績掲載','許可しない。許可する場合のみ、対象・掲載場所・時期を別途記入して合意。'),
      ('月額更新サポート','本契約には含まない。希望する場合は別の申込書で合意。'),
      ('個別条件','［なし／条件を明記］'),
    ]:field(d,label,value)
    d.add_heading('別紙2　確定見積書',1)
    d.add_paragraph('番号・版・発行日が一致する確定見積書PDFを本契約書と一緒に添付します。総額・工程内訳・支払期限・外部費用・振込先を記入したものを使用してください。')
    d.save(ROOT/'website-production-contract.docx')
    md=['# Webサイト制作契約書の条項ひな形','',f'版：{VERSION}。正式情報・案件仕様・見積書を記入してから相手方と合意します。','']
    for title,paragraphs in CONTRACT:
        md += ['## '+title,'']
        for p in paragraphs:md += [p,'']
    md += ['## 別紙・署名欄','', 'Word版に契約当事者の正式名称・氏名・住所・電話番号・署名欄と、別紙1（案件仕様）、別紙2（確定見積書の添付欄）があります。これらを本文と一緒に確定します。','']
    (ROOT/'contract-clauses.md').write_text('\n'.join(md),encoding='utf-8')

def write_support():
    d=document_base('更新サポート申込書・合意書','02 / OPTIONAL SUPPORT')
    for label,value in [
      ('申込番号・版','［案件番号 / 第1版］'),('発注者','［正式法人名・代表者／個人事業主の正式氏名］'),
      ('受注者','［運営者の正式氏名］（サービス名：専門サイト工房）'),('対象サイト','［URL］'),
      ('開始月・更新期間','［年・月］1日開始。以後毎月1日〜末日。'),('料金・上限','月5,500円（税込）。月2回・合計60分まで、繰越なし。'),
      ('依頼受付メール','eyu.eyu0806@gmail.com'),('発注者の担当・メール','［氏名・メール］'),
      ('開始月の支払期限','［年・月・日］'),('翌月以降の支払期限','月初。具体的期限は各月の請求書に記載。'),
      ('振込先','［銀行・支店・預金種別・口座番号・正式な口座名義］'),
      ('月額契約の解約窓口','当月末までに eyu.eyu0806@gmail.com へメール連絡、翌月から終了。')]:field(d,label,value)
    for title,paragraphs in SUPPORT:
        d.add_heading(title,1)
        for p in paragraphs:d.add_paragraph(p)
    field(d,'個別合意','［なし／途中月開始等の料金・期間・回数・支払日を明記］')
    d.add_heading('当事者・署名',1)
    for party in ['発注者','受注者']:
        field(d,party+'の正式情報','［正式氏名または法人名・代表者／住所／電話番号／メール］')
        field(d,party+'の承諾','［署名・電子署名／承諾メールの記録、日付］')
    d.save(ROOT/'optional-update-support.docx')

def build_estimate():
    wb=Workbook();wb.remove(wb.active)
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    thin=Side(style='thin',color=RULE)
    for name,base,scope in [('1ページ制作',55000,'1ページ・8セクションまで'),('小規模サイト制作',110000,'トップページを含む5ページまで')]:
        s=wb.create_sheet(name);s.sheet_view.showGridLines=False
        for col,width in {'A':5,'B':25,'C':19,'D':14,'E':14,'F':14,'G':18,'H':18}.items():s.column_dimensions[col].width=width
        def merge(row,left,right,text,fill=None,size=10,bold=False):
            s.merge_cells(start_row=row,start_column=left,end_row=row,end_column=right)
            c=s.cell(row,left,text);c.font=Font(name='Noto Sans CJK JP',size=size,bold=bold,color=INK)
            c.alignment=Alignment(vertical='center',wrap_text=True)
            if fill:c.fill=PatternFill('solid',fgColor=fill)
            return c
        def line(row,label,text,input=False,height=30):
            merge(row,1,2,label,PAPER,10,True)
            c=merge(row,3,8,text);s.row_dimensions[row].height=height
            if input:c.font=Font(name='Noto Sans CJK JP',size=10,color=GREEN)
        merge(1,1,8,'専門サイト工房 / pageatelier.jp',size=10);s.row_dimensions[1].height=22
        merge(2,1,8,'見積書｜'+name,size=23,bold=True);s.row_dimensions[2].height=40
        merge(3,1,8,'記入用ひな形：［ ］欄と案件条件を確定してから送付。緑文字の入力欄を編集してください。',PAPER);s.row_dimensions[3].height=27
        for r,label,text in [(5,'見積番号・版','［PA-2026-001 / 第1版］'),(6,'発行日・有効期限','［発行年月日］／［有効期限：発行日から14日等、具体的な年月日］'),(7,'宛先','［正式法人名・代表者／個人事業主の正式氏名］'),(8,'案件名・対象サイト','［案件名・URLまたは公開予定ドメイン］'),(9,'発行者（正式氏名）','［運営者の正式氏名］（個人事業・サービス名：専門サイト工房）'),(10,'発行者住所・電話','［郵便番号・正確な住所］／［連絡可能な電話番号］')]:line(r,label,text,True)
        line(11,'発行者メール','eyu.eyu0806@gmail.com')
        line(13,'基本範囲',scope+'。原稿支給のプライバシーポリシー1ページはページ数外で掲載。',height=38)
        line(14,'ページ・掲載内容','［別紙1と一致するページ／セクションの一覧・掲載量］',True,height=40)
        line(15,'含むもの','構成・デザイン・スマホ対応・公開設定・検索向け基本設定・フォーム1つ・地図・外部予約/SNSリンク・軽微な修正2回',height=45)
        line(16,'対象外・追加機能','CMS・決済・会員・検索・独自予約・新規原稿・撮影・大幅な変更等は別途見積り。追加対応は下表と契約別紙で確定。',height=44)
        for col,text in [(1,'No.'),(2,'制作項目'),(4,'数量'),(5,'単位'),(6,'税込単価'),(7,'税込金額')]:
            if col==2:merge(18,2,3,text,PAPER,10,True)
            elif col==7:merge(18,7,8,text,PAPER,10,True)
            else:merge(18,col,col,text,PAPER,10,True)
        s.row_dimensions[18].height=25
        rows=[(19,name,1,'式',base),(20,'ページ追加（既存デザイン）',0,'ページ',16500),(21,'追加修正（3回目以降）',0,'回',5500),(22,'個別見積りの追加機能',0,'式',0)]
        for r,item,qty,unit,price in rows:
            s.cell(r,1,r-18);merge(r,2,3,item);s.cell(r,4,qty);s.cell(r,5,unit);s.cell(r,6,price);merge(r,7,8,f'=D{r}*F{r}')
            s.row_dimensions[r].height=29
            for c in [4,6]:s.cell(r,c).font=Font(name='Noto Sans CJK JP',size=10,color=GREEN)
            s.cell(r,6).number_format='#,##0"円"';s.cell(r,7).number_format='#,##0"円"'
        dv=DataValidation(type='whole',operator='greaterThanOrEqual',formula1=0);dv.errorTitle='0以上の整数';dv.error='数量・金額は0以上の整数で記入してください。';dv.showErrorMessage=True;s.add_data_validation(dv)
        for address in ['D19:D22','F19:F22']:dv.add(address)
        merge(24,1,6,'制作費総額（税込）',PAPER,12,True);merge(24,7,8,'=SUM(G19:G22)',PAPER,15,True);s.row_dimensions[24].height=37
        merge(25,1,6,'うち消費税相当額（10％、総額から算出）');merge(25,7,8,'=ROUND(G24-G24/1.1,0)');s.row_dimensions[25].height=24
        merge(26,1,8,'追加単価は目安。掲載量・機能を確認して確定します。追加を含む場合は工程別内訳も調整。月額サポートは制作費に含みません。');s.row_dimensions[26].height=33
        merge(28,1,8,'工程別内訳（キャンセル精算の根拠）',size=12,bold=True);s.row_dimensions[28].height=25
        for r,label,ratio in [(29,'構成・設計',.2),(30,'デザイン',.3),(31,'実装・フォーム等の設定',.4),(32,'確認・公開・引渡し（残り）','=1-SUM(F29:F31)')]:
            merge(r,1,5,label);s.cell(r,6,ratio);s.cell(r,6).number_format='0%';s.cell(r,6).font=Font(name='Noto Sans CJK JP',size=10,color=GREEN if r<32 else INK)
            formula=f'=ROUND($G$24*F{r},0)' if r<32 else '=$G$24-SUM(G29:G31)'
            merge(r,7,8,formula);s.row_dimensions[r].height=25
        merge(33,1,5,'工程金額合計（制作費総額と一致）',PAPER);s.cell(33,6,'=SUM(F29:F32)');s.cell(33,6).number_format='0%';merge(33,7,8,'=SUM(G29:G32)',PAPER)
        merge(34,1,8,'配分は初期案。実際の作業に合わせて契約前に確定します。未完了の工程は作業記録・実施割合を示して精算し、工程に入っただけで全額請求しません。');s.row_dimensions[34].height=36
        merge(35,1,8,'=IF(F32<0,"工程配分エラー：構成・デザイン・実装の合計を100％以下にしてください","工程配分確認：合計100％。確認・公開の配分と金額は残りを自動計算します。")');s.row_dimensions[35].height=22
        ratios=DataValidation(type='custom',formula1='AND(F29>=0,F29<=1,SUM($F$29:$F$31)<=1)')
        ratios.showErrorMessage=True;ratios.errorTitle='工程配分を確認';ratios.error='0〜100％で、3工程の合計を100％以下にしてください。';s.add_data_validation(ratios);ratios.add('F29:F31')
        line(36,'着手前50％','=ROUND(G24/2,0)')
        # Keep amounts as real numeric formulas in G/H; payment dates stay editable.
        s.unmerge_cells('C36:H36');merge(36,3,6,'支払期限：［年・月・日］');merge(36,7,8,'=ROUND(G24/2,0)')
        line(37,'最終確認後・公開前50％','');s.unmerge_cells('C37:H37');merge(37,3,6,'支払期限：［年・月・日／最終承認後の請求日を確定］');merge(37,7,8,'=G24-G36')
        line(39,'制作開始・納期','［開始予定日］／［最終確認案日］／［公開予定日］。仕様・必要素材・着手金が揃ってから7日間目安。',True,height=42)
        merge(41,1,8,'見積条件・発注前の確認事項',size=17,bold=True);s.row_dimensions[41].height=36
        line(43,'原稿・素材・確認','原稿・写真は発注者支給。［支給期限・担当者］。完成案への回答は原則5営業日以内。無回答を自動承認とは扱いません。',True,height=48)
        line(44,'修正・追加作業','軽微な修正2回。まとめた要望への対応を1回と数え、制作ミス・仕様不一致の修正は回数外。追加作業は金額・範囲・納期を事前合意。',height=47)
        line(45,'フォーム・外部サービス','［送信項目・サービス・受信メール・受信承認担当］。制作体験で提案した機能はすべて基本料金に含むものではありません。',True,height=47)
        line(46,'対応環境・納品方法','［確認ブラウザ・幅／納品ZIP・リポジトリ・更新案内／公開アカウントの管理者］',True,height=43)
        line(47,'公開・権利','最終承認と残金入金後に公開・納品。新規成果物の権利移転・第三者ライセンス等は制作契約書で定めます。実績掲載は別途事前承諾。',height=47)
        line(49,'外部費用','制作費とは別。［ドメイン・サーバー・有料サービス等の項目／金額／負担者／支払先・期限／なし］',True,height=46)
        line(50,'銀行振込先','［銀行・支店・種別・口座番号・正式な口座名義］。振込手数料は発注者負担。',True,height=40)
        line(52,'契約前のキャンセル','無料相談・見積りの段階では費用は発生しません。')
        line(53,'契約後・着手前','受領金は原則返金。事前合意済みの外部実費が発生している場合は、その実費のみ精算。',height=40)
        line(54,'制作着手後','実施済みの工程・作業割合と事前合意済み外部実費を根拠に精算。未実施分を請求せず、余剰金は精算額合意後10営業日以内に返金。',height=48)
        line(55,'不具合・契約不履行','制作ミス・仕様不一致・制作側の不履行は、発注者都合のキャンセルとは区別し、契約と法令に従い修正・精算等を行います。',height=47)
        line(57,'任意の更新サポート','月5,500円（税込）、月2回・合計60分まで、繰越なし。別の申込書で合意し、本見積りの制作費には含めません。',height=42)
        line(58,'個別条件・契約添付','［なし／追加の合意事項］。この確定見積書PDFを、同じ番号・版の契約書に添付します。',True,height=44)
        merge(60,1,8,'この見積書だけで制作を開始しません。正式情報・案件仕様・納期・支払期限・外部費用を確認し、契約書と別紙一式への双方の承諾を保存します。',PAPER);s.row_dimensions[60].height=44
        merge(62,1,8,'書類ひな形 '+VERSION+'  /  専門サイト工房  /  https://pageatelier.jp/');s.row_dimensions[62].height=24
        for r in [24,25,29,30,31,32,33,36,37]:s.cell(r,7).number_format='#,##0"円"'
        for row in s:
            for c in row:
                if c.value is not None:
                    if c.font.name is None:c.font=Font(name='Noto Sans CJK JP',size=10,color=INK)
                    c.alignment=Alignment(vertical='center',wrap_text=True)
                    if c.row>=18 and c.row<=25 or c.row in [29,30,31,32,33,36,37]:c.border=Border(bottom=thin)
        s.freeze_panes='A5';s.print_area='A1:H62';s.sheet_properties.pageSetUpPr.fitToPage=True
        s.page_setup.orientation='portrait';s.page_setup.paperSize=s.PAPERSIZE_A4;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
        s.page_margins.left=.3;s.page_margins.right=.3;s.page_margins.top=.35;s.page_margins.bottom=.35
        s.row_breaks.append(Break(id=34));s.oddFooter.center.text='&P / &N';s.oddFooter.right.text='専門サイト工房';s.oddFooter.center.size=8
    wb.save(ROOT/'estimate-template.xlsx')

if __name__=='__main__':
    ROOT.mkdir(parents=True,exist_ok=True)
    write_contract();write_support();build_estimate()
    print('Generated blank contract, optional support agreement and two-plan estimate workbook in '+str(ROOT))
