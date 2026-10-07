"""Render the blank operation guides; never add private customer data here."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK

ROOT = Path(__file__).resolve().parent
INK, GREEN, MUTED, PAPER, RULE = '111111', '526657', '60605B', 'F4F4F0', 'C7C7BE'
FONT, DISPLAY = 'Noto Sans CJK JP', 'Noto Serif CJK JP'
LABELS = {
    '01-marketing-automation': ('01 / MARKETING', '記事・SNS・広告を、相談につなぐ。', '導入前の設計書。必要な接続・費用・公開条件を整理しています。'),
    '02-delivery-operations': ('02 / DELIVERY', '無料相談から、公開・引渡しまで。', '開始条件、2週間の工程、修正、承認、入金、納品を順番に進めます。'),
    '03-ai-prompts': ('03 / PROMPTS', '各工程で使う、35本のプロンプト。', '共通ルールと案件ブリーフを添えて、必要な工程だけコピーして使います。'),
    '04-acquisition-recommendation': ('04 / STRATEGY', '誰に、どこで、いくらで届けるか。', '業種・媒体・予算・接続の比較結果。外部サービスの最新料金・統計は未照合です。'),
    '05-launch-execution': ('05 / LAUNCH', '公開できたものと、接続するもの。', '実装・表示・公開の検証記録と、本人のログインが必要な接続の手順。'),
}

def set_font(run, name=FONT, size=None, color=None, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), name)
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None: run.bold = bold
    return run

def shade(element, fill):
    if hasattr(element, '_tc'): element=element._tc
    elif hasattr(element, '_p'): element=element._p
    props = element.get_or_add_tcPr() if element.tag == qn('w:tc') else element.get_or_add_pPr()
    node = OxmlElement('w:shd'); node.set(qn('w:fill'), fill); props.append(node)

def border_paragraph(p, color=RULE):
    borders = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    for k,v in [('val','single'),('sz','4'),('space','8'),('color',color)]:bottom.set(qn('w:'+k),v)
    borders.append(bottom);p._p.get_or_add_pPr().append(borders)

def field(p, name):
    r=OxmlElement('w:r'); f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),name)
    t=OxmlElement('w:t');t.text='1';r.append(t);f.append(r);p._p.append(f)

def inline(p, text, size=None, color=None):
    # Keep placeholders literal. Vertical bars escaped in Markdown tables are text.
    text=text.replace('\\|','|')
    for part in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)',text):
        if part.startswith('`') and part.endswith('`'):
            set_font(p.add_run(part[1:-1]),size=size,color=color or GREEN)
        elif part.startswith('**') and part.endswith('**'):
            set_font(p.add_run(part[2:-2]),size=size,color=color,bold=True)
        else:set_font(p.add_run(part),size=size,color=color)

def configure(d, label):
    sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    sec.header_distance=Cm(.85);sec.footer_distance=Cm(.85)
    normal=d.styles['Normal'];normal.font.name=FONT;normal.font.size=Pt(10)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
    normal.paragraph_format.line_spacing=1.25
    normal.paragraph_format.space_after=Pt(7)
    for name,size in [('Heading 1',19),('Heading 2',13),('Heading 3',11)]:
        s=d.styles[name];s.font.name=DISPLAY;s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string(INK)
        s._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),DISPLAY)
        s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(15);s.paragraph_format.space_after=Pt(8)
    hp=sec.header.paragraphs[0];set_font(hp.add_run('PAGE ATELIER  /  '+label),size=8,color=MUTED);border_paragraph(hp)
    fp=sec.footer.paragraphs[0];fp.paragraph_format.space_before=Pt(0)
    set_font(fp.add_run('2026-10-07  /  運用初期案     '),size=8,color=MUTED);field(fp,'PAGE')
    set_font(fp.add_run(' / '),size=8,color=MUTED);field(fp,'NUMPAGES')
    d.core_properties.author='Page Atelier';d.core_properties.subject='Business website production and marketing operations'

def cover(d,title,label,tagline,note):
    p=d.add_paragraph();p.paragraph_format.space_before=Pt(48)
    set_font(p.add_run(label),size=10,color=GREEN)
    p=d.add_paragraph();p.paragraph_format.space_before=Pt(24);p.paragraph_format.space_after=Pt(25)
    set_font(p.add_run(title),name=DISPLAY,size=27,color=INK)
    border_paragraph(p,GREEN)
    p=d.add_paragraph();p.paragraph_format.space_before=Pt(25)
    set_font(p.add_run(tagline),name=DISPLAY,size=17,color=INK)
    p=d.add_paragraph();inline(p,note,color=MUTED)
    p=d.add_paragraph();p.paragraph_format.space_before=Pt(45)
    set_font(p.add_run('専門サイト工房 / Page Atelier\n2026年10月7日版'),size=10,color=GREEN)
    p=d.add_paragraph();inline(p,'顧客情報・正式な運営者情報・認証情報は未記入です。記入済みの書類は非公開で保管してください。',size=9,color=MUTED)
    d.add_page_break()

def rows_from_table(lines):
    rows=[]
    for line in lines:
        if re.match(r'^\|\s*:?-',line):continue
        rows.append([x.strip().replace('\\|','|') for x in re.split(r'(?<!\\)\|',line)[1:-1]])
    return rows

def table(d, lines):
    rows=rows_from_table(lines);n=len(rows[0]);assert all(len(r)==n for r in rows)
    t=d.add_table(rows=0, cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    ratios={2:[.29,.71],3:[.20,.40,.40],4:[.15,.30,.27,.28]}.get(n,[1/n]*n)
    if rows[0][0]=='順':ratios=[.08,.48,.44]
    for col,ratio in zip(t.columns,ratios):col.width=Cm(17*ratio)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Cm(17*ratios[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cell,INK if i==0 else PAPER if i%2 else 'FFFFFF')
            p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.15
            inline(p,text,size=9,color='FFFFFF' if i==0 else INK)
            if i==0:
                for run in p.runs:run.bold=True
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                node=OxmlElement('w:'+side);node.set(qn('w:w'),'80');node.set(qn('w:type'),'dxa');margins.append(node)
            cell._tc.get_or_add_tcPr().append(margins)
        if i==0:
            repeat=OxmlElement('w:tblHeader');t.rows[i]._tr.get_or_add_trPr().append(repeat)
        # Keep each row readable on one page; table may split between rows.
        no_split=OxmlElement('w:cantSplit');t.rows[i]._tr.get_or_add_trPr().append(no_split)
    d.add_paragraph().paragraph_format.space_after=Pt(0)

def code(d, lines):
    for i,line in enumerate(lines):
        p=d.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.left_indent=Cm(.25);p.paragraph_format.right_indent=Cm(.25)
        p.paragraph_format.line_spacing=1.15
        p.paragraph_format.keep_with_next=i<len(lines)-1
        # Prompts and emails remain literal, selectable paragraphs in Word.
        shade(p,PAPER);set_font(p.add_run(line or ' '),size=9,color=INK)
    d.add_paragraph().paragraph_format.space_after=Pt(0)

def render(path):
    lines=path.read_text().splitlines();title=lines[0].removeprefix('# ')
    label,tagline,note=LABELS[path.stem];d=Document();configure(d,label);cover(d,title,label,tagline,note)
    p=d.add_paragraph();set_font(p.add_run('目次'),name=DISPLAY,size=18,color=INK)
    headings=[line[3:] for line in lines if line.startswith('## ')]
    for item in headings:
        p=d.add_paragraph();p.paragraph_format.space_after=Pt(5);inline(p,item,size=10)
    p=d.add_paragraph();inline(p,'Wordのナビゲーションウィンドウから各見出しへ移動できます。［ ］の項目は、確認した内容に置き換えてください。',size=9,color=MUTED)
    d.add_page_break()
    i=1
    while i<len(lines):
        line=lines[i]
        if not line.strip():i+=1;continue
        if line.startswith('```'):
            block=[];i+=1
            while i<len(lines) and not lines[i].startswith('```'):block.append(lines[i]);i+=1
            code(d,block);i+=1;continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].startswith('|'):block.append(lines[i]);i+=1
            table(d,block);continue
        if line.startswith('### '):d.add_paragraph(line[4:],style='Heading 2')
        elif line.startswith('## '):
            p=d.add_paragraph(line[3:],style='Heading 1')
            if path.stem=='03-ai-prompts' and line.startswith(('## 04','## 05','## 06','## 07','## 08')):p.paragraph_format.page_break_before=True
        else:
            p=d.add_paragraph();inline(p,line)
            if line.startswith('・'):p.paragraph_format.left_indent=Cm(.25)
        i+=1
    d.core_properties.title=title
    target=path.with_suffix('.docx');d.save(target);print(target.name)

if __name__=='__main__':
    for stem in LABELS:render(ROOT/(stem+'.md'))
